"""Underwriting model builder — Lev's headline agent deliverable ("provide a T-12 +
Rent Roll → build an Excel underwriting model: pro forma, DSCR, debt sizing"). Here
it's deterministic finance math (no LLM, no key) over a deal's fields, rendered to a
real .xlsx via openpyxl. `compute()` is pure + tested; `workbook()` just formats it.

The assumptions (amortization, NOI growth, hold, interest-only) are knobs, not
hardcodes — real underwriting needs to be tuned, so they're inputs with defaults."""
import io


def annual_debt_service(loan: float | None, rate_pct: float | None,
                        amort_years: int = 30, interest_only: bool = False) -> float | None:
    """Annual debt service on a fully-amortizing (or IO) loan. None if inputs missing.
    Guard the rate on presence (is None), not truthiness — a real 0% loan (seller
    financing / DPA) must reach the r==0 amortization branch, not return None."""
    if not loan or rate_pct is None:
        return None
    if interest_only:
        return loan * rate_pct / 100
    r = rate_pct / 100 / 12
    n = amort_years * 12
    pmt = loan / n if r == 0 else loan * r / (1 - (1 + r) ** (-n))
    return pmt * 12


def compute(deal: dict, amort_years: int = 30, noi_growth_pct: float = 3.0,
            hold_years: int = 5, interest_only: bool = False) -> dict:
    """Derive the underwriting metrics + a multi-year pro forma from a deal. Fills
    the gaps investors leave: NOI from price×cap, loan from price×LTV, and vice-versa."""
    price = deal.get("purchase_price")
    loan = deal.get("loan_amount")
    ltv = deal.get("ltv")
    rate = deal.get("interest_rate")
    cap = deal.get("cap_rate")
    noi = deal.get("noi")

    # fill in what's derivable
    if noi is None and price and cap:
        noi = round(price * cap / 100)
    if loan is None and price and ltv:
        loan = round(price * ltv / 100)
    if ltv is None and loan and price:
        ltv = round(loan / price * 100, 1)
    if cap is None and noi and price:
        cap = round(noi / price * 100, 2)

    ads = annual_debt_service(loan, rate, amort_years, interest_only)
    dscr = round(noi / ads, 2) if noi and ads else None
    debt_yield = round(noi / loan * 100, 2) if noi and loan else None
    # presence-guard so an all-cash deal (loan == 0) reports equity == full price
    equity = price - loan if price is not None and loan is not None else None
    # equity > 0 so an over-leveraged deal shows "—", not a sign-flipped positive return
    coc = round((noi - ads) / equity * 100, 2) if noi and ads and equity and equity > 0 else None

    projection = []
    if noi:
        for y in range(1, hold_years + 1):
            noi_y = round(noi * (1 + noi_growth_pct / 100) ** (y - 1))
            cf = round(noi_y - ads) if ads else None
            projection.append({
                "year": y, "noi": noi_y, "debt_service": round(ads) if ads else None,
                "cash_flow": cf, "dscr": round(noi_y / ads, 2) if ads else None})

    return {
        "assumptions": {"amort_years": amort_years, "noi_growth_pct": noi_growth_pct,
                        "hold_years": hold_years, "interest_only": interest_only},
        "purchase_price": price, "loan_amount": loan, "ltv": ltv, "interest_rate": rate,
        "cap_rate": cap, "noi": noi, "equity": equity,
        "annual_debt_service": round(ads) if ads else None,
        "dscr": dscr, "debt_yield_pct": debt_yield, "cash_on_cash_pct": coc,
        "projection": projection,
    }


def workbook(deal: dict, m: dict) -> bytes:
    """Render the computed model to an .xlsx (bytes). openpyxl only; no template file."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Underwriting"
    bold = Font(bold=True)

    def money(v):
        return f"${v:,.0f}" if isinstance(v, (int, float)) else "—"

    def pct(v):
        return f"{v}%" if isinstance(v, (int, float)) else "—"

    rows = [
        (deal.get("name", "Deal"), ""),
        (f"{deal.get('property_type') or '—'} · {deal.get('city') or ''} {deal.get('state') or ''}".strip(), ""),
        ("", ""),
        ("Sources & Uses", ""),
        ("Purchase price", money(m["purchase_price"])),
        ("Loan amount", money(m["loan_amount"])),
        ("Equity required", money(m["equity"])),
        ("LTV", pct(m["ltv"])),
        ("", ""),
        ("Debt sizing", ""),
        ("Interest rate", pct(m["interest_rate"])),
        ("Amortization (yrs)", "Interest-only" if m["assumptions"]["interest_only"] else m["assumptions"]["amort_years"]),
        ("Annual debt service", money(m["annual_debt_service"])),
        ("DSCR", m["dscr"] if m["dscr"] is not None else "—"),
        ("Debt yield", pct(m["debt_yield_pct"])),
        ("", ""),
        ("Returns", ""),
        ("Year-1 NOI", money(m["noi"])),
        ("Cap rate", pct(m["cap_rate"])),
        ("Cash-on-cash", pct(m["cash_on_cash_pct"])),
    ]
    for i, (a, b) in enumerate(rows, start=1):
        ws.cell(row=i, column=1, value=a)
        ws.cell(row=i, column=2, value=b)
        if a in ("Sources & Uses", "Debt sizing", "Returns") or (b == "" and a and i <= 2):
            ws.cell(row=i, column=1).font = bold

    if m["projection"]:
        base = len(rows) + 2
        ws.cell(row=base, column=1, value=f"Pro forma ({m['assumptions']['noi_growth_pct']}% NOI growth)").font = bold
        headers = ["Year", "NOI", "Debt service", "Cash flow", "DSCR"]
        for c, h in enumerate(headers, start=1):
            ws.cell(row=base + 1, column=c, value=h).font = bold
        for r, p in enumerate(m["projection"], start=base + 2):
            for c, key in enumerate(("year", "noi", "debt_service", "cash_flow", "dscr"), start=1):
                ws.cell(row=r, column=c, value=p[key])

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def demo() -> None:
    # 10M loan @ 6% / 30yr amortizing → ~719,463/yr debt service
    ads = annual_debt_service(10_000_000, 6.0, 30)
    assert abs(ads - 719_463) < 50, ads
    # interest-only is just rate × principal
    assert annual_debt_service(10_000_000, 6.0, 30, interest_only=True) == 600_000

    deal = {"name": "Harbor Pointe", "purchase_price": 41_000_000, "ltv": 68.0,
            "interest_rate": 6.5, "cap_rate": 5.2}
    m = compute(deal)
    assert m["loan_amount"] == round(41_000_000 * 0.68), m["loan_amount"]        # loan from LTV
    assert m["noi"] == round(41_000_000 * 0.052), m["noi"]                        # NOI from cap
    assert m["equity"] == 41_000_000 - m["loan_amount"], m["equity"]
    assert m["dscr"] and m["dscr"] == round(m["noi"] / m["annual_debt_service"], 2)
    assert len(m["projection"]) == 5 and m["projection"][0]["year"] == 1
    # NOI grows 3%/yr by default
    assert m["projection"][1]["noi"] > m["projection"][0]["noi"]
    # missing everything → no crash, all-None metrics
    empty = compute({"name": "TBD"})
    assert empty["dscr"] is None and empty["projection"] == []

    # 0% loan must amortize (annual = loan/amort_years), not return None
    assert abs(annual_debt_service(1_000_000, 0.0, 30) - 1_000_000 / 30) < 1
    assert annual_debt_service(1_000_000, 0.0, 30, interest_only=True) == 0.0
    # a missing (None) rate is still None
    assert annual_debt_service(1_000_000, None, 30) is None

    # all-cash deal (loan 0) → equity = full price, no phantom returns
    cash = compute({"purchase_price": 5_000_000, "loan_amount": 0, "noi": 300_000})
    assert cash["equity"] == 5_000_000 and cash["cash_on_cash_pct"] is None, cash

    # over-leveraged (equity < 0) must NOT report a positive cash-on-cash
    over = compute({"purchase_price": 10_000_000, "loan_amount": 12_000_000,
                    "noi": 800_000, "interest_rate": 6.0})
    assert over["equity"] == -2_000_000 and over["cash_on_cash_pct"] is None, over

    # workbook renders to real xlsx bytes (ZIP magic)
    assert workbook(deal, m)[:2] == b"PK"
    print("underwriting.demo OK")


if __name__ == "__main__":
    demo()
